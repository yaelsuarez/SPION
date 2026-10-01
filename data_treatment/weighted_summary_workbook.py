#!/usr/bin/env python3
"""Support-weighted summary of every metric, one row per classifier.

Reads Metrics_corrected_labels/ only; nothing is refitted or modified.

Weights are the per-class ground-truth voxel counts, so classes without Pill2
ground truth (0.05, 0.25) carry weight 0 and drop out on their own.

The macro (unweighted) average is written alongside, deliberately: tissue is
90.3% of the labelled voxels, so a support-weighted average is close to the
tissue value for every metric and hides the small classes entirely.
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from data_treatment.rank_corrected_labels import M, per_class_tables  # noqa: E402

TARGET = M / "summary.xlsx"
CLASSES = ["0", "0.1", "0.2", "0.3", "tissue"]
ORDER = ["logistic_regression", "random_forest", "gradient_boosting", "rbf_svm", "knn",
         "lda", "qda", "gmm", "hierarchical", "ordinal"]
FONT = "Arial"
METRICS = {
    "Precision": ("Precision", "higher is better"),
    "Recall": ("Recall", "higher is better"),
    "F1 / Dice": ("F1 (Dice)", "higher is better"),
    "AUC": ("AUC (one-vs-rest)", "higher is better"),
    "IoU": ("IoU (no threshold)", "higher is better"),
    "Brier score": ("Brier score", "lower is better"),
    "Voxel volume error (%)": (None, "lower is better; mean |signed relative %|"),
    "HD95 (mm)": ("HD95 (mm)", "lower is better"),
    "ASSD (mm)": ("ASSD (mm)", "lower is better"),
}


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    started = time.time()
    assert not TARGET.exists(), f"{TARGET.name} already exists; refusing to overwrite"

    tables = per_class_tables()
    vol = pd.read_csv(M / "Volume_error" / "volume_error.csv", dtype={"class": str})
    support = (vol[vol["class"].isin(CLASSES)]
               .pivot(index="classifier", columns="class", values="n_true_voxels")
               .reindex(ORDER)[CLASSES])
    abs_rel = (vol.assign(v=vol.relative_error_pct.abs())
               .pivot(index="classifier", columns="class", values="v")
               .reindex(ORDER)[CLASSES])

    per_class = {}
    for label, (key, _) in METRICS.items():
        per_class[label] = (abs_rel if key is None
                            else tables[key][0].reindex(ORDER).reindex(columns=CLASSES))

    def average(frame, weights=None):
        out = {}
        for clf in ORDER:
            v = frame.loc[clf].to_numpy(float)
            w = (np.ones_like(v) if weights is None
                 else support.loc[clf].to_numpy(float))
            m = np.isfinite(v)
            out[clf] = float(np.average(v[m], weights=w[m])) if m.any() else np.nan
        return pd.Series(out)

    weighted = pd.DataFrame({k: average(v, "support") for k, v in per_class.items()})
    macro = pd.DataFrame({k: average(v) for k, v in per_class.items()})

    # cross-check against sklearn's own weighted average
    checks = []
    for clf in ORDER:
        rep = pd.read_csv(M / "Confusion_matrix" / clf / "classification_report.csv",
                          index_col=0)
        row = rep.loc["weighted avg"]
        checks.append({"classifier": clf,
                       "precision_here": weighted.loc[clf, "Precision"],
                       "precision_sklearn": row["precision"],
                       "recall_here": weighted.loc[clf, "Recall"],
                       "recall_sklearn": row["recall"],
                       "f1_here": weighted.loc[clf, "F1 / Dice"],
                       "f1_sklearn": row["f1-score"]})
    check = pd.DataFrame(checks).set_index("classifier")
    for a, b in (("precision_here", "precision_sklearn"), ("recall_here", "recall_sklearn"),
                 ("f1_here", "f1_sklearn")):
        check[f"{a.split('_')[0]}_diff"] = (check[a] - check[b]).abs()
    worst = check[[c for c in check.columns if c.endswith("_diff")]].to_numpy().max()
    print(f"cross-check against sklearn 'weighted avg': max |difference| = {worst:.2e}")

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    body = Font(name=FONT)

    def sheet(name, frame, index_label, widths):
        ws = wb.create_sheet(name)
        ws.append([index_label] + [str(c) for c in frame.columns])
        for key, row in frame.iterrows():
            ws.append([key] + [None if pd.isna(v) else float(v) for v in row])
        for cell in ws[1]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body
        ws.freeze_panes = "B2"
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    widths = [24] + [16] * len(METRICS)
    sheet("Weighted average", weighted.round(6), "classifier", widths)
    sheet("Macro average", macro.round(6), "classifier", widths)
    sheet("Class support (weights)", support, "classifier", [24] + [13] * len(CLASSES))
    for label in METRICS:
        name = label.replace("/", "-").replace("(", "").replace(")", "")[:31]
        sheet(name, per_class[label].round(6), "classifier", [24] + [13] * len(CLASSES))

    ws = wb.create_sheet("_README")
    rows = [
        ("Content", "support-weighted average of each metric across classes, per classifier"),
        ("Weights", "per-class ground-truth voxel count; classes without ground truth "
                    "(0.05, 0.25) carry weight 0 and drop out"),
        ("Classes averaged", ", ".join(CLASSES)),
        ("CAUTION", "tissue is 347,114 of 384,315 voxels (90.3%), so the weighted average "
                    "is close to the tissue value for every metric. The macro average is "
                    "provided alongside for that reason."),
        ("Cross-check", f"weighted precision/recall/F1 match sklearn's 'weighted avg' row "
                        f"to {worst:.1e}"),
        ("Labels", "CORRECTED: Pill2 Segmentation_4_0.25 treated as 0.2"),
        ("Source", "MRI/Classifiers/Metrics_corrected_labels/"),
        ("Recomputed", "no; averages of stored per-class values, nothing refitted"),
        ("Created", time.strftime("%Y-%m-%d %H:%M:%S")),
    ] + [(f"Direction — {k}", v[1]) for k, v in METRICS.items()]
    for key, value in rows:
        ws.append([key, value])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = body
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 88

    wb.save(TARGET)
    pd.set_option("display.width", 230)
    print("\nWEIGHTED AVERAGE")
    print(weighted.round(3).to_string())
    print("\nMACRO AVERAGE (unweighted, for comparison)")
    print(macro.round(3).to_string())
    print(f"\nsaved {TARGET}\nelapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
