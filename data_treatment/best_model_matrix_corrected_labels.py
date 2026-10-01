#!/usr/bin/env python3
"""One table: rows are metrics, columns are classes, cells name the best model.

Built by reading Metrics_corrected_labels/ only. Nothing is refitted or
modified; a single new workbook is written.

Columns cover the five classes with Pill2 ground truth under the corrected
labelling. 0.05 and 0.25 are omitted: several metrics are undefined for them,
and the ones that are defined measure false confidence rather than performance,
so a "best model" there would not mean what the column header implies.
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402

from data_treatment.rank_corrected_labels import M, ROOT, per_class_tables, rank  # noqa: E402

TARGET = ROOT / "best_model_per_metric_and_class_corrected_labels.xlsx"
CLASSES = ["0", "0.1", "0.2", "0.3", "tissue"]
FONT = "Arial"


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    started = time.time()
    assert not TARGET.exists(), f"{TARGET.name} already exists; refusing to overwrite"

    tables = per_class_tables()

    # the signed variant, ranked by closeness to zero rather than by magnitude
    signed = pd.read_csv(M / "Volume_error" /
                         "volume_error_signed_relative_pct_by_classifier.csv"
                         ).set_index("classifier")
    signed.columns = [str(c) for c in signed.columns]

    rows, sources = [], []
    for label, (frame, direction, fmt, source) in tables.items():
        arrow = "higher is better" if direction == "max" else "lower is better"
        entry = {"Metric": f"{label} — {arrow}"}
        for c in CLASSES:
            entry[c] = rank(frame[c], direction, fmt)[0] if c in frame.columns else "–"
        rows.append(entry)
        sources.append(f"{label}: {source}")

    entry = {"Metric": "Volume error, signed relative (%) — closest to zero is best"}
    for c in CLASSES:
        s = signed[c].dropna()
        ordered = s.reindex(s.abs().sort_values().index)
        entry[c] = f"{ordered.index[0]} ({ordered.iloc[0]:+.1f})" if len(ordered) else "–"
    rows.append(entry)
    sources.append("Volume error, signed relative (%): Metrics_corrected_labels/Volume_error/"
                   "volume_error_signed_relative_pct_by_classifier.csv")

    table = pd.DataFrame(rows)[["Metric"] + CLASSES]

    wb = Workbook()
    ws = wb.active
    ws.title = "Best model per metric"
    ws.append(list(table.columns))
    for record in table.itertuples(index=False, name=None):
        ws.append(list(record))

    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 52
    for i in range(2, len(CLASSES) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 30

    foot = ws.max_row + 2
    note = Font(name=FONT, italic=True, size=9)
    ws.cell(foot, 1, "Each cell names the best-performing classifier for that metric and "
                     "class, with its value. CORRECTED LABELS: Pill2 Segmentation_4_0.25 "
                     "is treated as 0.2.").font = note
    ws.cell(foot + 1, 1, "Ranking direction is stated in each metric name. Distance metrics "
                         "use 0.32 x 0.32 x 0.32 mm isotropic spacing. Values transcribed "
                         "from the stored metric files; nothing was refitted.").font = note
    ws.cell(foot + 2, 1, "Classes 0.05 and 0.25 are omitted: they have no Pill2 ground truth "
                         "under the corrected labelling.").font = note
    for j, line in enumerate(sources, start=foot + 4):
        ws.cell(j, 1, f"Source — {line}").font = note

    wb.save(TARGET)
    pd.set_option("display.width", 250)
    pd.set_option("display.max_colwidth", 30)
    print(table.to_string(index=False))
    print(f"\n{len(table)} metrics x {len(CLASSES)} classes -> {TARGET}")
    print(f"elapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
