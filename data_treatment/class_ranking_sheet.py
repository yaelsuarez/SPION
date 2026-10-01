#!/usr/bin/env python3
"""Rank the classifiers on one class alone, as a new sheet in summary.xlsx.

Usage: class_ranking_sheet.py [class]   (default 0.2)

Values come from the per-metric sheets already in the workbook, so the ranking
is consistent with every other sheet there. Nothing is recomputed and no model
is refitted.
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402

from data_treatment.macro_ranking_sheets import DIRECTION, TARGET, ranks_of  # noqa: E402

FONT = "Arial"
#: display label -> the sheet holding that metric's per-class values
SHEETS = {"Precision": "Precision", "Recall": "Recall", "F1 / Dice": "F1 - Dice",
          "AUC": "AUC", "IoU": "IoU", "Brier score": "Brier score",
          "Voxel volume error (%)": "Voxel volume error %", "HD95 (mm)": "HD95 mm",
          "ASSD (mm)": "ASSD mm"}


def main(target_class: str = "0.2") -> int:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    started = time.time()
    name = f"Ranking class {target_class}"
    book = pd.ExcelFile(TARGET)
    assert name not in book.sheet_names, f"{name} already exists"

    values = {}
    for label, sheet in SHEETS.items():
        frame = pd.read_excel(book, sheet).set_index("classifier")
        frame.columns = [str(c) for c in frame.columns]
        assert target_class in frame.columns, f"{sheet} has no column {target_class}"
        values[label] = frame[target_class]
    values = pd.DataFrame(values)
    ranked = ranks_of(values).sort_values("mean_rank")
    values = values.reindex(ranked.index)

    wb = load_workbook(TARGET)
    ws = wb.create_sheet(name)
    ws.append(["classifier"] + [str(c) for c in ranked.columns]
              + [f"{c} value" for c in values.columns])
    for key in ranked.index:
        ws.append([key]
                  + [float(v) if isinstance(v, float) else int(v) for v in ranked.loc[key]]
                  + [None if pd.isna(v) else float(v) for v in values.loc[key]])
    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    for cell in ws[1]:
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    readme = wb["_README"]
    readme.append([name, f"classifiers ranked on class {target_class} alone, 1 = best per "
                         "metric; rank columns first, then the underlying values"])
    for cell in readme[readme.max_row]:
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    wb.save(TARGET)

    assert all(sorted(ranked[m]) == list(range(1, 11)) for m in DIRECTION)
    pd.set_option("display.width", 240)
    print(f"RANKING ON CLASS {target_class} ALONE (1 = best)\n")
    print(ranked.to_string())
    print("\nunderlying values\n")
    print(values.round(4).to_string())
    print(f"\nsaved sheet '{name}' to {TARGET}\nelapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1 else "0.2"))
