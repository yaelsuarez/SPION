#!/usr/bin/env python3
"""Organise the finished analysis into three Excel workbooks.

Nothing is recomputed, refitted, retrained or modified. Every value is read from
a stored analysis output and written out unchanged.

Source choices worth knowing, both recorded on each workbook's _README sheet:

* Raw rs vs I0 comes from Normalized_data_i0_rs/Segmentations/, whose I0_raw is
  the fitted amplitude and so is independent of any normalisation.
* Normalised rs vs I0_normalized uses the FINAL label-light scheme:
  Segmentations_training/ for the three training phantoms and the complete Pill2
  store. The I0_normalized column in Normalized_data_i0_rs/Segmentations/ is the
  superseded Strategy-B version (factors 6452.85 / 1229.50 / 1304.21 / 2011.36)
  and is deliberately not used.

Large sheets are streamed with openpyxl's write_only mode; pandas.to_excel holds
every cell in memory and is far too slow at this size.
"""

import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import DATA_ROOT  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

MRI = DATA_ROOT
RAW_STORE = MRI / "Normalized_data_i0_rs" / "Segmentations"
LABEL_LIGHT = MRI / "Normalized_data_i0_rs" / "Segmentations_training"
PILL2_COMPLETE = MRI / "i0_rs" / "Volume_for_testing" / "Pill2_segmentation_complete" / "voxels_complete.csv.gz"
PILL2_META = MRI / "i0_rs" / "Volume_for_testing" / "Pill2_normalized" / "metadata.json"
DIAG = MRI / "Training" / "diagnosis"
OUT = MRI / "Results" / "rs_vs_i0"
SAMPLES = ["Syringes", "3D", "Pill1", "Pill2"]
FONT = "Arial"
EXCEL_MAX_ROWS = 1_048_576


def write_workbook(path, sheets, readme):
    """Stream frames to sheets; readme is a list of (key, value) rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    for name, frame in sheets.items():
        assert len(frame) < EXCEL_MAX_ROWS, f"{name}: {len(frame):,} rows exceeds the Excel limit"
        ws = wb.create_sheet(name[:31])
        ws.freeze_panes = "A2"
        from openpyxl.cell import WriteOnlyCell
        header = [str(c) for c in frame.columns]
        bold = Font(name=FONT, bold=True)
        head = []
        for text in header:
            cell = WriteOnlyCell(ws, value=text)
            cell.font = bold
            head.append(cell)
        ws.append(head)
        for record in frame.itertuples(index=False, name=None):
            ws.append([None if isinstance(v, float) and np.isnan(v) else v for v in record])
        for i, width in enumerate([max(11, min(22, len(h) + 2)) for h in header], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
    ws = wb.create_sheet("_README")
    for key, value in readme:
        ws.append([key, value])
    wb.save(path)
    return path


def raw_and_normalized() -> tuple[dict, dict]:
    raw_sheets, norm_sheets = {}, {}
    pill2_factor = json.loads(PILL2_META.read_text())["normalization"]["factor"]

    for sample in SAMPLES:
        parts = []
        for seg in sorted(p for p in (RAW_STORE / sample).iterdir() if p.is_dir()):
            f = pd.read_csv(seg / "voxels.csv.gz", dtype={"material": str})
            f["material"] = seg.name.split("_", 2)[2]     # folder name is authoritative
            parts.append(f)
        raw = pd.concat(parts, ignore_index=True); del parts
        raw_sheets[sample] = raw[["sample", "segmentation", "material", "z", "y", "x",
                                  "rs_selected", "I0_raw", "B_selected", "model_selected",
                                  "n_informative", "flag_high_rs", "flag_low_information",
                                  "quality_ok"]]
        print(f"  raw        {sample:9s} {len(raw):>8,} voxels, "
              f"{raw.material.nunique()} classes")

        if sample == "Pill2":
            p = pd.read_csv(PILL2_COMPLETE, dtype={"material": str})
            p["material"] = p.segmentation.str.split("_", n=2).str[2]
            p["sample"] = "Pill2"
            p["normalization_factor"] = pill2_factor
            norm = p[["sample", "segmentation", "material", "z", "y", "x", "rs_selected",
                      "I0_selected", "I0_normalized", "I0_residual", "normalization_factor",
                      "origin", "flag_high_rs", "flag_low_information", "quality_ok"]]
        else:
            parts = []
            for seg in sorted(p for p in (LABEL_LIGHT / sample).iterdir() if p.is_dir()):
                f = pd.read_csv(seg / "voxels.csv.gz", dtype={"material": str})
                f["material"] = seg.name.split("_", 2)[2]
                parts.append(f)
            norm = pd.concat(parts, ignore_index=True); del parts
            norm = norm[["sample", "segmentation", "material", "z", "y", "x", "rs_selected",
                         "I0_selected", "I0_normalized", "I0_residual",
                         "normalization_factor", "flag_high_rs", "flag_low_information",
                         "quality_ok"]]
        norm_sheets[sample] = norm
        print(f"  normalized {sample:9s} {len(norm):>8,} voxels, "
              f"factor {norm.normalization_factor.iloc[0]:g}")
    return raw_sheets, norm_sheets


def failed_parameters() -> dict:
    """Long-format evidence per tested feature, from the diagnostic outputs."""
    rows = []

    def push(feature, analysis, source, dataset, group, statistic, value, n=None):
        rows.append({"feature": feature, "analysis": analysis, "dataset": dataset,
                     "class_or_pair": group, "statistic": statistic, "value": value,
                     "n": n, "source_file": source})

    def melt_stats(path, feature_col, stats):
        d = pd.read_csv(path, dtype=str)
        src = str(path.relative_to(MRI))
        for _, r in d.iterrows():
            for s in stats:
                if s in d.columns and pd.notna(r[s]):
                    push(r[feature_col], "distribution by class and phantom", src,
                         r.get("sample", ""), r.get("material", ""), s,
                         float(r[s]), r.get("n"))

    melt_stats(DIAG / "shape_parameters" / "stats_by_class_phantom.csv", "candidate",
               ["median", "q1", "q3", "p05", "p95"])
    melt_stats(DIAG / "derived_features" / "feature_stats_by_class_phantom.csv", "feature",
               ["median", "q1", "q3", "p05", "p95"])

    for path, col in ((DIAG / "shape_parameters" / "phantom_invariance_pairs.csv", "candidate"),
                      (DIAG / "derived_features" / "same_class_cross_phantom.csv", "feature")):
        d = pd.read_csv(path, dtype=str)
        src = str(path.relative_to(MRI))
        for _, r in d.iterrows():
            push(r[col], "same class, different phantom (higher = more invariant)", src,
                 f"{r.sample_a} vs {r.sample_b}", r.material, "overlap", float(r.overlap))

    d = pd.read_csv(DIAG / "shape_parameters" / "separation_within_and_across_phantoms.csv",
                    dtype=str)
    src = str((DIAG / "shape_parameters" / "separation_within_and_across_phantoms.csv")
              .relative_to(MRI))
    for _, r in d.iterrows():
        push(r.candidate, f"class separation, {r.kind} (lower = better separated)", src,
             f"{r.sample_a} vs {r.sample_b}", r.pair, "overlap", float(r.overlap))

    d = pd.read_csv(DIAG / "derived_features" / "class_pair_separation.csv", dtype=str)
    src = str((DIAG / "derived_features" / "class_pair_separation.csv").relative_to(MRI))
    for _, r in d.iterrows():
        if pd.notna(r.overlap):
            push(r.feature, "class separation (lower = better separated)", src,
                 r.dataset, r.pair, "overlap", float(r.overlap))

    d = pd.read_csv(DIAG / "shape_parameters" / "candidate_summary.csv", dtype=str)
    src = str((DIAG / "shape_parameters" / "candidate_summary.csv").relative_to(MRI))
    for _, r in d.iterrows():
        for c in d.columns:
            if c in ("candidate", "new") or pd.isna(r[c]):
                continue
            push(r.candidate, "summary", src, "", "", c, float(r[c]))

    # ---- B_selected and B_normalized -------------------------------------
    d = pd.read_csv(DIAG / "B_parameter" / "B_by_class.csv", dtype=str)
    src = str((DIAG / "B_parameter" / "B_by_class.csv").relative_to(MRI))
    for _, r in d.iterrows():
        for s in ("median_B", "q1", "q3", "pct_B_zero"):
            if s in d.columns and pd.notna(r[s]):
                push("B_selected", "distribution by class", src, r.dataset, r.material,
                     s, float(r[s]), r.get("n"))

    d = pd.read_csv(DIAG / "B_parameter" / "overlap_with_without_B.csv", dtype=str)
    src = str((DIAG / "B_parameter" / "overlap_with_without_B.csv").relative_to(MRI))
    for _, r in d.iterrows():
        for c in ("rs", "rs + I0_residual", "rs + B", "rs + I0_residual + B"):
            if c in d.columns and pd.notna(r[c]):
                push("B_selected", "feature-set overlap (lower = better separated)", src,
                     r.dataset, f"{r.class_a} vs {r.class_b}", c, float(r[c]))

    nt = DIAG / "B_parameter" / "normalization_transfer"
    for filename, analysis in (("class_phantom_statistics.csv", "distribution by class and phantom"),
                               ("same_class_cross_phantom_overlap.csv",
                                "same class, different phantom (higher = more invariant)"),
                               ("class_pair_separation.csv",
                                "class separation (lower = better separated)"),
                               ("training_vs_pill2_transfer.csv", "training vs Pill2 transfer")):
        path = nt / filename
        if not path.is_file():
            continue
        d = pd.read_csv(path, dtype=str)
        src = str(path.relative_to(MRI))
        for _, r in d.iterrows():
            for c in d.columns:
                if pd.isna(r[c]):
                    continue
                try:
                    value = float(r[c])
                except ValueError:
                    continue
                push("B_normalized", analysis, src,
                     r.get("sample", r.get("dataset", r.get("sample_a", ""))),
                     r.get("material", r.get("pair", "")), c, value, r.get("n"))

    # ---- dAIC as a classifier feature -------------------------------------
    d = pd.read_csv(DIAG / "dAIC" / "model_comparison.csv", dtype=str)
    src = str((DIAG / "dAIC" / "model_comparison.csv").relative_to(MRI))
    for _, r in d.iterrows():
        for c in ("balanced_accuracy", "macro_f1", "accuracy"):
            if c in d.columns and pd.notna(r[c]):
                push("dAIC", "classifier evaluation on Pill2", src, r.feature_set, "",
                     c, float(r[c]), r.get("n_voxels"))
    for tag in ("rs", "rs_I0_residual", "rs_dAIC", "rs_I0_residual_dAIC"):
        path = DIAG / "dAIC" / f"per_class_metrics_{tag}.csv"
        if not path.is_file():
            continue
        d = pd.read_csv(path, index_col=0)
        src = str(path.relative_to(MRI))
        for cls, r in d.iterrows():
            if cls in ("accuracy", "macro avg", "weighted avg") or r.get("support", 0) == 0:
                continue
            for c in ("precision", "recall", "f1-score"):
                push("dAIC", "classifier per-class metrics on Pill2", src, tag, str(cls),
                     c, float(r[c]), int(r["support"]))

    everything = pd.DataFrame(rows)
    sheets = {}
    for feature, block in everything.groupby("feature"):
        name = str(feature).replace("/", "_")[:31]
        sheets[name] = block.reset_index(drop=True)
    return sheets


def main() -> int:
    started = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    for filename in ("rs_vs_i0.xlsx", "rs_vs_i0_normalized.xlsx", "failed_parameters.xlsx"):
        assert not (OUT / filename).exists(), f"{filename} already exists; refusing to overwrite"

    raw_sheets, norm_sheets = raw_and_normalized()
    write_workbook(OUT / "rs_vs_i0.xlsx", raw_sheets, [
        ("Content", "Raw rs vs I0 per voxel, one sheet per sample, all classes"),
        ("Source", "MRI/Normalized_data_i0_rs/Segmentations/<sample>/<segmentation>/voxels.csv.gz"),
        ("I0 column", "I0_raw - the fitted amplitude, independent of any normalisation"),
        ("rs column", "rs_selected - the AIC-selected decay rate, 1/ms"),
        ("material", "taken from the segmentation folder name, not the CSV column"),
        ("Recomputed", "no; values transcribed unchanged"),
        ("Created", time.strftime("%Y-%m-%d %H:%M:%S")),
    ])
    print(f"  wrote rs_vs_i0.xlsx  ({(OUT/'rs_vs_i0.xlsx').stat().st_size/1e6:.0f} MB)")

    factors = {s: float(f.normalization_factor.iloc[0]) for s, f in norm_sheets.items()}
    write_workbook(OUT / "rs_vs_i0_normalized.xlsx", norm_sheets, [
        ("Content", "rs vs I0_normalized per voxel, one sheet per sample, all classes"),
        ("Normalisation", "final label-light scheme"),
        ("Factors", ", ".join(f"{s} {v:g}" for s, v in factors.items())),
        ("Source, training", "MRI/Normalized_data_i0_rs/Segmentations_training/"),
        ("Source, Pill2", "MRI/i0_rs/Volume_for_testing/Pill2_segmentation_complete/"),
        ("Not used", "Normalized_data_i0_rs/Segmentations I0_normalized - superseded "
                     "Strategy-B factors 6452.85 / 1229.50 / 1304.21 / 2011.36"),
        ("I0_residual", "I0_normalized minus f(rs), the degree-2 polynomial fitted on "
                        "Syringes + 3D + Pill1 only"),
        ("Recomputed", "no; values transcribed unchanged"),
        ("Created", time.strftime("%Y-%m-%d %H:%M:%S")),
    ])
    print(f"  wrote rs_vs_i0_normalized.xlsx  "
          f"({(OUT/'rs_vs_i0_normalized.xlsx').stat().st_size/1e6:.0f} MB)")

    failed = failed_parameters()
    write_workbook(OUT / "failed_parameters.xlsx", failed, [
        ("Content", "one sheet per tested feature that did not prove useful"),
        ("Layout", "long format: feature, analysis, dataset, class_or_pair, statistic, "
                   "value, n, source_file"),
        ("Overlap direction", "phantom invariance: higher = distributions agree. "
                              "class separation: lower = better separated."),
        ("Sources", "MRI/Training/diagnosis/{shape_parameters, derived_features, "
                    "B_parameter, dAIC}"),
        ("Recomputed", "no; values transcribed unchanged"),
        ("Created", time.strftime("%Y-%m-%d %H:%M:%S")),
    ])
    print(f"  wrote failed_parameters.xlsx  "
          f"({(OUT/'failed_parameters.xlsx').stat().st_size/1e6:.1f} MB)")
    for name, frame in failed.items():
        print(f"      {name:16s} {len(frame):>5,} rows")
    print(f"\nsaved to {OUT}\nelapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
