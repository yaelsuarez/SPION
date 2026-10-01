#!/usr/bin/env python3
"""Rank-agreement (bump chart) data: how each classifier ranks under each metric.

Reads Metrics_corrected_labels/ only. Nothing is refitted or modified.

Each metric is reduced to one number per classifier - the mean over the five
classes that have Pill2 ground truth - then classifiers are ranked 1 (best) to
10 under that metric's own direction. A bump chart plots rank against metric,
one line per classifier.
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402

from data_treatment.rank_corrected_labels import M, per_class_tables  # noqa: E402

OUT = M / "bump_chart"
TARGET = OUT / "rank_agreement_bump_chart.xlsx"
CLASSES = ["0", "0.1", "0.2", "0.3", "tissue"]
FONT = "Arial"
#: label -> (table key, direction, note). Matches the manuscript figure panels.
METRICS = {
    "Precision": ("Precision", "max", "higher is better"),
    "Recall": ("Recall", "max", "higher is better"),
    "F1 (Dice)": ("F1 (Dice)", "max", "higher is better"),
    "AUC": ("AUC (one-vs-rest)", "max", "higher is better"),
    "ASSD (mm)": ("ASSD (mm)", "min", "lower is better"),
    "Brier": ("Brier score", "min", "lower is better"),
    "Volume error (%)": (None, "min", "mean |signed relative %|, closest to zero is best"),
}


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    started = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    assert not TARGET.exists(), f"{TARGET.name} already exists; refusing to overwrite"

    tables = per_class_tables()
    signed = pd.read_csv(M / "Volume_error" /
                         "volume_error_signed_relative_pct_by_classifier.csv"
                         ).set_index("classifier")
    signed.columns = [str(c) for c in signed.columns]

    values = {}
    for label, (key, direction, _) in METRICS.items():
        if key is None:
            values[label] = signed[CLASSES].abs().mean(axis=1)
        else:
            values[label] = tables[key][0].reindex(columns=CLASSES).mean(axis=1)
    values = pd.DataFrame(values)

    ranks = pd.DataFrame({
        label: values[label].rank(ascending=(METRICS[label][1] == "min")).astype(int)
        for label in METRICS})
    ranks["best_rank"] = ranks[list(METRICS)].min(axis=1)
    ranks["worst_rank"] = ranks[list(METRICS)].max(axis=1)
    ranks["rank_spread"] = ranks.worst_rank - ranks.best_rank
    ranks["times_ranked_first"] = ranks[list(METRICS)].eq(1).sum(axis=1)
    ranks = ranks.sort_values(["best_rank", "rank_spread"])
    values = values.reindex(ranks.index)

    long = [{"classifier": c, "metric": m, "rank": int(ranks.loc[c, m]),
             "value": float(values.loc[c, m]), "direction": METRICS[m][1]}
            for c in ranks.index for m in METRICS]
    long = pd.DataFrame(long)

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    body = Font(name=FONT)

    def sheet(name, frame, index_label=None, widths=None):
        ws = wb.create_sheet(name)
        ws.append(([index_label] if index_label else []) + [str(c) for c in frame.columns])
        for key, row in frame.iterrows():
            ws.append(([key] if index_label else []) +
                      [None if pd.isna(v) else (int(v) if isinstance(v, (int,)) else float(v))
                       for v in row])
        for cell in ws[1]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body
        ws.freeze_panes = "B2"
        for i, w in enumerate(widths or ([24] + [15] * frame.shape[1]), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    sheet("Ranks", ranks, "classifier")
    sheet("Values", values.round(6), "classifier")
    ws = wb.create_sheet("Long")
    ws.append(list(long.columns))
    for record in long.itertuples(index=False, name=None):
        ws.append(list(record))
    for cell in ws[1]:
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
    ws.freeze_panes = "A2"
    for i, w in enumerate([24, 20, 10, 16, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("_README")
    for key, value in [
        ("Purpose", "bump chart: rank (y) against metric (x), one line per classifier"),
        ("Sheets", "Ranks = 1 (best) to 10; Values = the macro value ranked; "
                   "Long = tidy format for plotting"),
        ("Macro definition", f"mean over the classes with Pill2 ground truth: "
                             f"{', '.join(CLASSES)}"),
        ("Labels", "CORRECTED: Pill2 Segmentation_4_0.25 treated as 0.2"),
        ("Source", "MRI/Classifiers/Metrics_corrected_labels/"),
        ("Recomputed", "no; ranks derived from stored metric values, nothing refitted"),
        ("Created", time.strftime("%Y-%m-%d %H:%M:%S")),
    ]:
        ws.append([key, value])
    for label, (_, direction, note) in METRICS.items():
        ws.append([f"Direction — {label}", f"{direction} ({note})"])
    for cell_row in ws.iter_rows():
        for cell in cell_row:
            cell.font = body
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 78

    wb.save(TARGET)
    pd.set_option("display.width", 200)
    print(ranks.to_string())
    print(f"\nsaved {TARGET}")
    print(f"elapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
