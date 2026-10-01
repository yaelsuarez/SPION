#!/usr/bin/env python3
"""Top-3 ranking and the per-class workbook, from the corrected-label metrics.

Reads Classifiers/Metrics_corrected_labels/ only. Nothing is refitted, retrained
or recomputed - every value is transcribed from a stored metric file. Two new
files are written; no existing file is touched.

Sheets follow the corrected labelling: 0.25 no longer has Pill2 ground truth and
0.2 does, so the per-class workbook covers 0, 0.1, 0.2, 0.3 and tissue.
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import DATA_ROOT  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

ROOT = DATA_ROOT / "Classifiers"
M = ROOT / "Metrics_corrected_labels"
CSV_TARGET = ROOT / "metric_top3_models_corrected_labels.csv"
XLSX_TARGET = ROOT / "best_models_by_class_corrected_labels.xlsx"
ORDER = ["logistic_regression", "random_forest", "gradient_boosting", "rbf_svm", "knn",
         "lda", "qda", "gmm", "hierarchical", "ordinal"]
CLASSES = ["0", "0.05", "0.1", "0.2", "0.25", "0.3", "tissue"]
KEY_CLASSES = ["0", "0.1", "0.2", "0.25", "0.3", "tissue"]
SHEETS = ["0", "0.1", "0.2", "0.3", "tissue"]     # classes with Pill2 ground truth
THRESHOLDS = [("no threshold", "all_probability"), ("p>=0.5", "0.5_probability"),
              ("p>=0.7", "0.7_probability"), ("p>=0.75", "0.75_probability"),
              ("p>=0.8", "0.8_probability")]
FONT = "Arial"


def long_table(path, value, dtype_class=True):
    d = pd.read_csv(path, dtype={"class": str} if dtype_class else None)
    return d.pivot(index="classifier", columns="class", values=value)


def per_class_tables():
    """{label: (frame indexed by classifier, direction, fmt, source)}."""
    tables = {}
    for label, folder, filename in (("Precision", "Precision", "precision_summary.csv"),
                                    ("Recall", "Recall", "recall_summary.csv"),
                                    ("F1 (Dice)", "F1_Dice", "f1_dice_summary.csv")):
        t = pd.read_csv(M / folder / filename, index_col=0)
        t.columns = [str(c) for c in t.columns]
        tables[label] = (t, "max", "{:.4f}", f"Metrics_corrected_labels/{folder}/{filename}")

    auc = pd.read_csv(M / "AUC" / "auc_summary.csv").set_index("classifier")
    wide = auc[[c for c in auc.columns if c.startswith("auc_class_")]]
    wide.columns = [c.replace("auc_class_", "") for c in wide.columns]
    tables["AUC (one-vs-rest)"] = (wide, "max", "{:.4f}",
                                   "Metrics_corrected_labels/AUC/auc_summary.csv")

    for tag, folder in THRESHOLDS:
        path = M / "IoU" / folder / "iou_summary.csv"
        if not path.is_file():
            continue
        t = pd.read_csv(path).set_index("classifier")
        t = t[[f"IoU_{c}" for c in CLASSES if f"IoU_{c}" in t.columns]]
        t.columns = [c.replace("IoU_", "") for c in t.columns]
        tables[f"IoU ({tag})"] = (t, "max", "{:.4f}",
                                  f"Metrics_corrected_labels/IoU/{folder}/iou_summary.csv")

    tables["MCC (one-vs-rest)"] = (long_table(M / "MCC" / "mcc.csv", "mcc_ovr"), "max",
                                   "{:.4f}", "Metrics_corrected_labels/MCC/mcc.csv")
    tables["Brier score"] = (long_table(M / "Brier" / "brier.csv", "brier"), "min",
                             "{:.4f}", "Metrics_corrected_labels/Brier/brier.csv")

    vol = pd.read_csv(M / "Volume_error" / "volume_error.csv", dtype={"class": str})
    vol["abs_rel"] = vol.relative_error_pct.abs()
    tables["Volume error, absolute relative (%)"] = (
        vol.pivot(index="classifier", columns="class", values="abs_rel"), "min", "{:.1f}",
        "Metrics_corrected_labels/Volume_error/volume_error.csv")
    tables["Volume error, absolute (mm3)"] = (
        vol.pivot(index="classifier", columns="class", values="absolute_error_mm3"),
        "min", "{:.1f}", "Metrics_corrected_labels/Volume_error/volume_error.csv")

    tables["HD95 (mm)"] = (long_table(M / "HD95" / "hd95.csv", "hd95_mm"), "min", "{:.2f}",
                           "Metrics_corrected_labels/HD95/hd95.csv")
    tables["ASSD (mm)"] = (long_table(M / "ASSD" / "assd.csv", "assd_mm"), "min", "{:.2f}",
                           "Metrics_corrected_labels/ASSD/assd.csv")
    return tables


def rank(series, direction, fmt):
    s = pd.Series(series).astype(float)
    s = s[np.isfinite(s)]
    if s.empty:
        return ["not rankable: no values available", "", ""]
    if s.nunique() == 1:
        return [f"all {len(s)} models tie ({fmt.format(s.iloc[0])})", "", ""]
    ordered = s.sort_values(ascending=(direction == "min"))
    return [f"{ordered.index[i]} ({fmt.format(ordered.iloc[i])})" if i < len(ordered) else ""
            for i in range(3)]


def build_csv(tables):
    summary = pd.read_csv(M / "summary.csv").set_index("classifier")
    rows = []

    def add(parameter, series, direction, fmt="{:.4f}"):
        best = rank(series, direction, fmt)
        rows.append({"Parameter": parameter, "Best model 1": best[0],
                     "Best model 2": best[1], "Best model 3": best[2]})

    add("Balanced_accuracy_overall", summary["balanced_accuracy"], "max")
    add("Macro_F1_overall", summary["macro_f1"], "max")
    add("Confusion_matrix_overall_accuracy", summary["accuracy"], "max")
    add("AUC_macro", summary["auc_macro"], "max")
    add("AUC_weighted", summary["auc_weighted"], "max")
    add("MCC_multiclass", summary["mcc_multiclass"], "max")
    add("MCC_ovr_macro", summary["mcc_ovr_macro"], "max")
    add("Brier_multiclass", summary["brier_multiclass"], "min")
    add("Brier_macro", summary["brier_macro"], "min")
    add("HD95_macro_mm", summary["hd95_macro_mm"], "min", "{:.2f}")
    add("ASSD_macro_mm", summary["assd_macro_mm"], "min", "{:.2f}")
    add("Volume_error_mean_abs_relative", summary["mean_abs_relative_volume_error"], "min")

    for tag, folder in THRESHOLDS:
        path = M / "IoU" / folder / "iou_summary.csv"
        if not path.is_file():
            continue
        t = pd.read_csv(path).set_index("classifier")
        add(f"IoU_macro_at_{tag}", t["IoU_macro_mean"], "max")
        add(f"Coverage_at_{tag}", t["coverage_pct"], "max", "{:.2f}")

    for label, (frame, direction, fmt, _) in tables.items():
        for c in KEY_CLASSES:
            if c in frame.columns:
                add(f"{label}_{c}", frame[c], direction, fmt)

    out = pd.DataFrame(rows)[["Parameter", "Best model 1", "Best model 2", "Best model 3"]]
    out.to_csv(CSV_TARGET, index=False)
    return out


def build_xlsx(tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(name=FONT, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    body = Font(name=FONT)
    note = Font(name=FONT, italic=True, size=9)

    for cls in SHEETS:
        ws = wb.create_sheet(cls)
        ws.append(["Metric", "Best model", "2nd best model", "3rd best model"])
        sources = []
        for label, (frame, direction, fmt, source) in tables.items():
            if cls not in frame.columns:
                continue
            arrow = "higher is better" if direction == "max" else "lower is better"
            ws.append([f"{label} — {arrow}"] + rank(frame[cls], direction, fmt))
            sources.append(f"{label}: {source}")
        for cell in ws[1]:
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = body
                cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for i, width in enumerate([44, 34, 34, 34], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        foot = ws.max_row + 2
        ws.cell(foot, 1, f"Class '{cls}'. CORRECTED LABELS: Pill2 Segmentation_4_0.25 is "
                         "treated as 0.2. Values transcribed from the stored metric files; "
                         "nothing was refitted or recomputed.").font = note
        ws.cell(foot + 1, 1, "Ranking direction is stated in each metric name. Distance "
                             "metrics use 0.32 x 0.32 x 0.32 mm isotropic spacing.").font = note
        for j, line in enumerate(sources, start=foot + 3):
            ws.cell(j, 1, f"Source — {line}").font = note
    wb.save(XLSX_TARGET)


def main() -> int:
    started = time.time()
    for path in (CSV_TARGET, XLSX_TARGET):
        assert not path.exists(), f"{path.name} already exists; refusing to overwrite"

    tables = per_class_tables()
    out = build_csv(tables)
    build_xlsx(tables)
    pd.set_option("display.width", 220)
    pd.set_option("display.max_colwidth", 34)
    print(out.to_string(index=False))
    print(f"\n{len(out)} rows -> {CSV_TARGET.name}")
    print(f"{len(SHEETS)} sheets -> {XLSX_TARGET.name}")
    print(f"elapsed {time.time()-started:.0f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
